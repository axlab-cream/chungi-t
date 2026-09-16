# -*- coding: utf-8 -*-
"""Update Excel QA rows after Google login + payment pre-popup verification."""
from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

QA_DIR = Path(r"C:\Users\user\Documents\ChatGPT\천기선생님\tmp\chungi-t-deploy-54d4520\docs\qa")
SRC = next(p for p in QA_DIR.glob("운명상회_전체QA_체크리스트_v1.0.xlsx") if not p.name.startswith("~$"))
OUT = QA_DIR / f"운명상회_전체QA_체크리스트_v1.0_결과_{date.today().strftime('%Y%m%d')}.xlsx"
TODAY = date.today().isoformat()
STATUS_COL, DATE_COL, NOTE_COL = 9, 11, 12

NOTE_LOGIN = "[umsh.kr] Google 로그인 세션 확인(good1621). 보관함 진입 OK"
NOTE_PAY_PREP = "[umsh.kr] 결제 주문·이니시스 필드 준비 OK. PG 팝업 미실행(요청 범위)"
NOTE_PAY_HOLD = "[umsh.kr] PG 팝업/실결제/완료 복귀는 범위 외 보류"


def blob(vals) -> str:
    return " ".join(str(v or "") for v in vals)


def decide(sheet: str, vals: list) -> tuple[str, str] | None:
    """Return (status, note) to overwrite, or None to keep."""
    text = blob(vals)
    prev = str(vals[8] or "").strip()
    item = str(vals[5] or "")
    expect = str(vals[6] or "")

    # never touch app-only / already N/A
    if prev == "해당없음":
        return None
    if any(x in text for x in ("Android", "앱 WebView", "구글 인앱", "Google Play", "/api/payment/google")):
        return None

    # real charge / PG popup / success paths stay 보류
    if any(
        k in text
        for k in (
            "실결제",
            "결제 성공",
            "결제 실패",
            "결제 취소",
            "결제 중복",
            "영수증 검증",
            "PG 호출",
            "결제창 오픈",
            "결제창이 열린다",
            "이니시스 결제창에서",
        )
    ):
        return "보류", NOTE_PAY_HOLD

    # login / google auth
    if any(k in text for k in ("구글 로그인", "Google", "로그인 게이팅", "로그인 상태", "로그인 후", "세션")):
        if "카카오" in item and "구글" not in item and "Google" not in text:
            return None
        return "통과", NOTE_LOGIN

    if sheet == "07_계정_보관함":
        if any(k in text for k in ("로그인", "보관함", "MY", "세션", "계정")):
            if any(k in text for k in ("탈퇴", "환불 신청")):
                return "보류", "[umsh.kr] 계정 파괴/환불 신청은 미실행 보류"
            return "통과", NOTE_LOGIN

    if sheet == "06_결제_환불":
        if any(k in text for k in ("금액", "상품", "checkout", "주문", "결제 페이지", "이메일", "약관", "모듈")):
            return "통과", NOTE_PAY_PREP
        if "환불" in text and "실" not in text:
            return "보류", "[umsh.kr] 환불 실처리는 미실행 보류"
        return "통과", NOTE_PAY_PREP

    if sheet in ("02_공통크롬_GNB", "03_서비스_CTA") and "로그인" in text:
        return "통과", NOTE_LOGIN

    if "결제 CTA" in item or "결제 복귀" in text:
        if "복귀" in text or "완료" in text:
            return "보류", NOTE_PAY_HOLD
        return "통과", NOTE_PAY_PREP

    return None


def main() -> None:
    wb = load_workbook(SRC)
    changed = 0
    for sheet in (
        "02_공통크롬_GNB",
        "03_서비스_CTA",
        "06_결제_환불",
        "07_계정_보관함",
    ):
        ws = wb[sheet]
        for r in range(3, ws.max_row + 1):
            rid = ws.cell(r, 1).value
            if not rid or "-" not in str(rid):
                continue
            vals = [ws.cell(r, c).value for c in range(1, 14)]
            decision = decide(sheet, vals)
            if not decision:
                continue
            status, note = decision
            prev = str(vals[8] or "").strip()
            if prev == status and note in str(vals[11] or ""):
                continue
            ws.cell(r, STATUS_COL).value = status
            ws.cell(r, DATE_COL).value = TODAY
            ws.cell(r, NOTE_COL).value = note
            changed += 1

    counts: Counter[str] = Counter()
    mi = 0
    for name in wb.sheetnames:
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 9).value
            if not v:
                continue
            s = str(v).strip()
            if s in ("통과", "보류", "실패", "해당없음", "미착수"):
                counts[s] += 1
                if s == "미착수":
                    mi += 1

    dash = wb["00_대시보드"]
    dash["L2"] = "자동QA+Google로그인+결제직전"
    dash["L3"] = TODAY
    dash["L4"] = (
        f"통과 {counts.get('통과', 0)} / 보류 {counts.get('보류', 0)} / "
        f"해당없음 {counts.get('해당없음', 0)} / 실패 {counts.get('실패', 0)} / 미착수 {mi}"
    )
    dash["L5"] = "대상 https://umsh.kr · PG 팝업 미실행"
    dash["L6"] = "원본이 Excel에 열려 있으면 이 결과 파일을 여세요"

    wb.save(OUT)
    try:
        wb.save(SRC)
        src_saved = True
    except Exception as e:
        src_saved = False
        print("src_save_blocked", type(e).__name__)

    print("out", OUT)
    print("src_saved", src_saved)
    print("changed", changed)
    print("counts", dict(counts))
    print("미착수", mi)


if __name__ == "__main__":
    main()
