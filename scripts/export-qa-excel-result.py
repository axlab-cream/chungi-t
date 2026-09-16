# -*- coding: utf-8 -*-
"""Copy QA workbook to an unlocked result file (when Excel holds a lock)."""
from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

QA_DIR = Path(r"C:\Users\user\Documents\ChatGPT\천기선생님\tmp\chungi-t-deploy-54d4520\docs\qa")
SRC = next(p for p in QA_DIR.glob("*.xlsx") if not p.name.startswith("~$") and "결과" not in p.name)
OUT = QA_DIR / f"운명상회_전체QA_체크리스트_v1.0_결과_{date.today().strftime('%Y%m%d')}.xlsx"


def main() -> None:
    wb = load_workbook(SRC)
    counts: Counter[str] = Counter()
    for name in wb.sheetnames:
        ws = wb[name]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 9).value
            if v and str(v).strip() in ("통과", "보류", "실패", "해당없음", "미착수"):
                counts[str(v).strip()] += 1

    dash = wb["00_대시보드"]
    dash["L2"] = "자동QA 반영"
    dash["L3"] = date.today().isoformat()
    dash["L4"] = (
        f"통과 {counts.get('통과', 0)} / 보류 {counts.get('보류', 0)} / "
        f"해당없음 {counts.get('해당없음', 0)} / 실패 {counts.get('실패', 0)} / "
        f"미착수 {counts.get('미착수', 0)}"
    )
    dash["L5"] = "대상: https://umsh.kr"
    dash["L6"] = "원본이 Excel에서 열려 있으면 이 결과 파일을 여세요. 대시보드 수치는 Excel에서 수식 재계산됩니다."

    wb.save(OUT)
    print("src", SRC)
    print("out", OUT)
    print("counts", dict(counts))


if __name__ == "__main__":
    main()
